import re
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def parse_and_add_markdown(doc, md_text):
    """
    Parses basic Markdown (bold, lists, tables) and adds them to a python-docx Document.
    """
    lines = md_text.split("\n")
    in_table = False
    table_data = []

    for line in lines:
        line = line.strip()
        if not line:
            if in_table:
                # Render table
                _render_table(doc, table_data)
                in_table = False
                table_data = []
            continue

        # Check for tables
        if line.startswith("|") and line.endswith("|"):
            in_table = True
            # Ignore markdown table separator lines like |---|---|
            if set(line.replace("|", "").replace("-", "").replace(" ", "").replace(":", "")) == set():
                continue
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            table_data.append(cells)
            continue
        
        if in_table:
            _render_table(doc, table_data)
            in_table = False
            table_data = []

        # Check for lists
        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            _parse_inline_markdown(p, line[2:])
            continue
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            _parse_inline_markdown(p, re.sub(r'^\d+\.\s', '', line))
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        _parse_inline_markdown(p, line)

    if in_table:
        _render_table(doc, table_data)

def _parse_inline_markdown(paragraph, text):
    """
    Parses **bold** and *italic* within a string and appends them as runs to the paragraph.
    """
    # Find all **bold** and *italic*
    # We use a simple tokenization strategy
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for token in tokens:
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)

def _render_table(doc, table_data):
    if not table_data:
        return
    
    rows = len(table_data)
    cols = max(len(row) for row in table_data)
    
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Table Grid'
    
    for i, row_data in enumerate(table_data):
        row_cells = table.rows[i].cells
        for j, cell_text in enumerate(row_data):
            if j < len(row_cells):
                # Apply inline formatting to table cells
                p = row_cells[j].paragraphs[0]
                _parse_inline_markdown(p, cell_text)
                
                # Make header bold
                if i == 0:
                    for run in p.runs:
                        run.bold = True
